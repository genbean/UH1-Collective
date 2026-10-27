# Build HeliSR_Rev4.0_Board_BOM.xlsx: only the parts that go on the board (the Pro Micro and its pin headers in a
# section of their own), checked against the footprints in the board file. One input cell sets how many boards to
# build; quantities and prices follow from it.
# usage: make_board_bom_v4.py <HeliSR_Rev4.0.kicad_pcb> <out.xlsx>
import re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

board, out = sys.argv[1:3]
DK = "https://www.digikey.com/en/products/result?keywords="
PRICES = ("Unit prices for the parts carried over from Rev 3.0 are single-unit Digi-Key prices from the Rev 3.0 cart "
          "(2026-10-19 to 2026-10-22). The new Rev 4.0 parts are priced from distributor listings found on 2026-10-26 "
          "(PCA9546ADR: Digi-Key cut tape; BM04B-SRSS-TBT: Digi-Key via an aggregator) or, for the 0805 resistors and "
          "capacitor, at the Rev 3.0 0805 resistor's price: check them in the cart. Before shipping and tariff. The bare "
          "PCB is priced at a third of an OSH Park order of three (Rev 3.0: $10.20).")

# (refs, value, footprint, description, manufacturer, mfr part, digikey part, unit USD, notes)
SOLDERED = [
    ("U1, U2", "74HC165", "SOIC-16_3.9x9.9mm_P1.27mm", "Shift register, 74HC165, SOIC-16", "Texas Instruments",
     "SN74HC165DR", "296-8250-1-ND", 0.50, "Front, under the Pro Micro: fit before the sockets"),
    ("RN1, RN2", "10k x8 bussed", "R_Array_SIP9", "Resistor network, 8 x 10 kOhm bussed, SIP-9", "Bourns",
     "4609X-101-103LF", "4609X-101-103LF-ND", 0.39, "Pin 1 (dot) is the common pin: it goes in the square pad. Keep upright"),
    ("C1, C2", "100nF", "C_1206_3216Metric", "Capacitor, 100 nF, 50 V, X7R, 1206", "Samsung", "CL31B104KBCNNNC",
     "1276-1017-1-ND", 0.14, "Decoupling: C1 beside U1's VCC pin, C2 beside U2's"),
    ("C3", "22uF", "C_1206_3216Metric", "Capacitor, 22 uF, 10 V, X5R, 1206", "Samsung", "CL31A226KPHNNNE",
     "1276-1287-1-ND", 0.47, "Bulk supply capacitor, between U1 and U2"),
    ("U3", "PCA9546A", "SOIC-16_3.9x9.9mm_P1.27mm", "I2C switch, 4 channels, with reset, SOIC-16 (narrow, D package)",
     "Texas Instruments", "PCA9546ADR", "296-21770-1-ND", 1.82,
     "Back. Gives each AS5600 port its own I2C channel (they all answer at 0x36). Not the wide DW package"),
    ("C4", "100nF", "C_0805_2012Metric", "Capacitor, 100 nF, 50 V, X7R, 0805", "Samsung", "CL21B104KBCNNNC",
     "1276-1003-1-ND", 0.10, "Back, beside U3's VCC pin"),
    ("R1-R8", "2.2k", "R_0805_2012Metric", "Resistor, 2.2 kOhm, 1 %, 1/8 W, 0805", "YAGEO", "RC0805FR-072K2L",
     "311-2.20KCRCT-ND", 0.10, "Back: the ports' I2C pull-ups, a pair under each port (R1/R2 AX1, R3/R4 AX2, R5/R6 AX3, "
     "R7/R8 AX4; the odd one SDA, the even one SCL)"),
    ("R9-R11", "10k", "R_0805_2012Metric", "Resistor, 10 kOhm, 1 %, 1/8 W, 0805", "YAGEO", "RC0805FR-0710KL",
     "311-10.0KCRCT-ND", 0.10, "Back: pull-ups on the Pro Micro's SDA (R9) and SCL (R10), and on U3's /RESET (R11)"),
    ("SW1", "TL1015AF160QG", "SW_TL1015AF160QG", "Tactile switch, SMD, 3.9 x 2.9 mm, gull wing", "E-Switch",
     "TL1015AF160QG", "EG4344CT-ND", 0.48, "RST, at the far end between AX4 and AX2. The four ground tabs are left unsoldered"),
    ("J1, J2", None, "JST_SH_BM10B-SRSS-TB_1x10-1MP_P1.00mm_Vertical", "JST SR/SH header, 10-pin, 1.0 mm, SMD, top entry",
     "JST", "BM10B-SRSS-TBT", "455-BM10B-SRSS-TBTCT-ND", 1.19,
     "Switch inputs: J1 IN1-IN8, J2 IN9-IN16. Digi-Key estimates $0.33 tariff each on top"),
    ("J11-J14", None, "JST_SH_BM04B-SRSS-TB_1x04-1MP_P1.00mm_Vertical", "JST SR/SH header, 4-pin, 1.0 mm, SMD, top entry",
     "JST", "BM04B-SRSS-TBT", "455-BM04B-SRSS-TBTCT-ND", 0.71,
     "AS5600 ports AX1-AX4 (J11-J14): 1 SDA, 2 GND, 3 VCC, 4 SCL. Fit all four, or only the ports you use"),
    ("J7, J8", None, "PinSocket_1x12_P2.54mm_Vertical", "Socket header, female, 1x12, 2.54 mm, through hole", "Sullins",
     "PPTC121LFBN-RC", "S6100-ND", 0.75, "Pro Micro sockets. Fit last"),
]
PLUGGED = [
    ("Module", 1, "Pro Micro 5V/16 MHz, ATmega32U4", "SparkFun", "12640", "1568-12640-ND", 23.99,
     "Runs the HeliSR firmware (Rev 4.0 build). Plugs into J7/J8"),
    ("Module", 2, "Pin header, male, 1x12, 2.54 mm", "Sullins", "PRPC012SAAN-RC", "35-PRPC012SAAN-RC-ND", 0.24,
     "Soldered to the Pro Micro"),
]

# ---- check the list against the board: every placed part once, with the right footprint and value ------------
def footprints(path):
    s = open(path, encoding="utf-8").read()
    out, i = {}, 0
    while (i := s.find("\n\t(footprint ", i)) >= 0:
        d, j, q = 0, i + 1, False
        while True:                                  # find the matching paren, skipping quoted strings
            ch = s[j]
            if q:
                if ch == "\\": j += 1
                elif ch == '"': q = False
            elif ch == '"': q = True
            elif ch == "(": d += 1
            elif ch == ")":
                d -= 1
                if d == 0: break
            j += 1
        fp = s[i + 2:j + 1]; i = j                    # from "(footprint" to its closing paren
        ref = re.search(r'\(property "Reference" "([^"]*)"', fp).group(1)
        if "(exclude_from_bom" in fp or re.search(r"\(attr [^)]*exclude_from_bom", fp): continue
        out[ref] = (re.match(r'\(footprint "(?:[^":]*:)?([^"]+)"', fp).group(1),
                    re.search(r'\(property "Value" "([^"]*)"', fp).group(1))
    return out
def expand(refs):
    r = []
    for part in refs.split(", "):
        m = re.fullmatch(r"([A-Z]+)(\d+)-[A-Z]*(\d+)", part)
        r += [f"{m.group(1)}{k}" for k in range(int(m.group(2)), int(m.group(3)) + 1)] if m else [part]
    return r
fps = footprints(board)
listed = [ref for line in SOLDERED for ref in expand(line[0])]
assert sorted(listed) == sorted(fps) and len(set(listed)) == len(listed), (sorted(set(fps) ^ set(listed)), listed)
for refs, value, fpname, *_ in SOLDERED:
    for ref in expand(refs):
        assert fps[ref][0] == fpname, (ref, fps[ref], fpname)
        assert value is None or fps[ref][1] == value, (ref, fps[ref], value)
print(f"{len(fps)} parts on the board, all listed once with matching footprints")

# ---- workbook ---------------------------------------------------------------------------------------------------
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap, centre = Alignment(wrap_text=True, vertical="top"), Alignment(horizontal="center", vertical="top")
USD = '"$"#,##0.00;-"$"#,##0.00;"-"'
HEAD = ["#", "Reference", "Qty per board", "Qty to order", "Value", "Description", "Footprint", "Manufacturer",
        "Mfr part number", "Digi-Key part number", "Unit price", "Ext. price", "Notes"]
WIDTH = [4, 12, 8, 8, 13, 36, 30, 15, 18, 25, 9, 10, 46]
wb = Workbook(); ws = wb.active; ws.title = "Board BOM"
wb.calculation = CalcProperties(fullCalcOnLoad=True)
ws["A1"] = "HeliSR Rev 4.0 - board BOM"; ws["A1"].font = F(bold=True, size=14)
ws["A2"] = ("The parts you solder onto the bare HeliSR Rev 4.0 PCB, checked against the board file. Rev 4.0 has not "
            "been ordered. The front carries the connectors, the button and Rev 3.0's parts; the back carries U3, C4 "
            "and R1-R11. The Pro Micro and its pin headers are listed separately at the end. " + PRICES)
ws["A2"].font = F(italic=True, color="595959"); ws["A2"].alignment = wrap
ws.merge_cells("A2:M2"); ws.row_dimensions[2].height = 42
ws["A3"] = "Boards to build"; ws["A3"].font = F(bold=True); ws.merge_cells("A3:B3")
ws["C3"] = 1; ws["C3"].font = F(bold=True, color="0000FF"); ws["C3"].fill = PatternFill("solid", fgColor="FFFF00")
ws["C3"].alignment = centre; ws["C3"].border = box
ws["D3"] = "Edit the yellow cell. Quantities to order and prices follow from it."
ws["D3"].font = F(italic=True, color="595959"); ws.merge_cells("D3:M3")
HR = 5
for c, (h, w) in enumerate(zip(HEAD, WIDTH), 1):
    cell = ws.cell(HR, c, h); cell.font = F(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F3864"); cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = box; ws.column_dimensions[cell.column_letter].width = w
ws.row_dimensions[HR].height = 28
r, n = HR, 0
def section(title):
    global r
    r += 1
    ws.cell(r, 1, title).font = F(bold=True)
    for c in range(1, len(HEAD) + 1):
        ws.cell(r, c).fill = PatternFill("solid", fgColor="D9E1F2"); ws.cell(r, c).border = box
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
def line(refs, per, value, desc, fpname, mfr, mpn, dk, unit, notes):
    global r, n
    r += 1; n += 1
    vals = [n, refs, per, f"=C{r}*$C$3", value, desc, fpname, mfr, mpn, dk, unit, f"=IF(K{r}=\"\",\"\",D{r}*K{r})", notes]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v); cell.font = F(); cell.alignment = wrap; cell.border = box
    for c in (1, 3, 4): ws.cell(r, c).alignment = centre
    for c in (11, 12): ws.cell(r, c).number_format = USD
    if dk:
        ws.cell(r, 10).hyperlink = DK + dk; ws.cell(r, 10).font = F(color="0563C1", underline="single")
    return r
section("Soldered to the board")
first = r + 1
for refs, value, fpname, desc, mfr, mpn, dk, unit, notes in SOLDERED:
    vals = sorted({fps[x][1] for x in expand(refs)})
    line(refs, len(expand(refs)), value or ", ".join(vals), desc, fpname, mfr, mpn, dk, unit, notes)
line("PCB", 1, "HeliSR Rev 4.0", "Bare PCB (no parts fitted), 2 layers, 1.6 mm FR4, 39.3 x 33.64 mm", "", "OSH Park",
     "", "", 3.40, "Not ordered yet. Priced as Rev 3.0: $10.20 for three bare boards from OSH Park")
last = r
r += 1
ws.cell(r, 11, "Subtotal").font = F(bold=True)
ws.cell(r, 12, f"=SUM(L{first}:L{last})").font = F(bold=True); ws.cell(r, 12).number_format = USD
sub1 = r
section("Plugs into the board (the Pro Micro and its headers)")
first2 = r + 1
for refs, per, desc, mfr, mpn, dk, unit, notes in PLUGGED:
    line(refs, per, "", desc, "", mfr, mpn, dk, unit, notes)
last2 = r
r += 1
ws.cell(r, 11, "Subtotal").font = F(bold=True)
ws.cell(r, 12, f"=SUM(L{first2}:L{last2})").font = F(bold=True); ws.cell(r, 12).number_format = USD
r += 1
ws.cell(r, 11, "Total").font = F(bold=True)
ws.cell(r, 12, f"=L{sub1}+L{r - 1}").font = F(bold=True); ws.cell(r, 12).number_format = USD
r += 1
ws.cell(r, 1, "The bare PCB counts as a third of an OSH Park order of three ($10.20 for Rev 3.0).").font = F(italic=True, color="595959")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEAD))
ws.freeze_panes = f"A{HR + 1}"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = f"{HR}:{HR}"
wb.save(out)
print("wrote", out, "with", n, "lines")
